import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation('original_backup/Growth Plan 500 Registrations in 7 Days.pptx')

# Update Slide 5 content to reflect the upgraded working product
slide5 = prs.slides[4]
for shape in slide5.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if 'not yet connected to a real form' in text:
            shape.text_frame.text = 'Full-stack referral tracking, live student registration gateway, and automated pacing engine. Production-grade demonstrable application with end-to-end attribution.'
        elif 'https://claude.ai/artifact/' in text or 'Demo:' in text:
            shape.text_frame.text = 'Live Local Demo: http://localhost:3000/ | Standalone HTML: index.html (Zero-dependency Node server & REST API)'

# Professional Speaker Notes for Interview Presentation
speaker_notes = [
    # Slide 1 Notes
    """SLIDE 1: TARGET AUDIENCE & VALUE PROPOSITION
Key Talking Points:
1. Target Persona: Final-year B.Tech/BE students (CSE, IT, ECE, AI/DS) facing campus placements, particularly across tier-2 and tier-3 colleges.
2. The Core Bottleneck: Students have AI coursework on their resumes, but zero deployed projects that impress technical interviewers.
3. The High-Converting Promise: Not a generic webinar or passive lecture. In exactly 60 minutes, students build, test, and deploy a real working AI project to put directly on their GitHub and resume.
4. Growth Mindset: Everything presented here is a hypothesis designed for rapid validation within the first 48 hours.""",

    # Slide 2 Notes
    """SLIDE 2: THREE-CHANNEL DISTRIBUTION STRATEGY
Key Talking Points:
1. Channel A - WhatsApp Connectors (Target: 250 regs): Peer distribution beats brand posts. 10 campus connectors post vetted copy into placement groups. Connector incentive: ₹50 thank-you per connector once they reach 15+ registrations (performance-gated).
2. Channel B - Clubs & Communities (Target: 140 regs): Coding clubs and student chapters receive a customized 30-second live build teaser.
3. Channel C - Referral Viral Loop (Target: 117 regs): Registrants receive an immediate invite link on the confirmation screen to unlock an AI Placement Toolkit. 25% share rate × 1.2 sign-ups = K-factor of 0.30.
4. Capital Efficiency: ₹800 (40% of total budget) is held strictly in reserve as a conditional paid test only if Day 4 pacing lags.""",

    # Slide 3 Notes
    """SLIDE 3: 7-DAY MILESTONE SCHEDULE & ₹2,000 BUDGET
Key Talking Points:
1. Cumulative Milestone Curve: 40 (D1) -> 120 (D2) -> 210 (D3) -> 300 (D4) -> 390 (D5) -> 450 (D6) -> 500 (D7).
2. The Critical Day 4 Checkpoint: Day 4 target is 300 cumulative registrations. 90% pacing threshold is 270. If we are under 270 on Day 4, contingency triggers fire immediately.
3. Strict Budget Allocation: ₹500 for connectors, ₹600 for referral leaderboard prizes (₹300/₹200/₹100), ₹800 paid reserve, ₹100 payment buffer = ₹2,000 total.
4. Unit Economics: Effective CAC at 500 registrations is exactly ₹4.00 per student.""",

    # Slide 4 Notes
    """SLIDE 4: FUNNEL ARITHMETIC & CONTINGENCY ENGINE
Key Talking Points:
1. Funnel Mathematics:
   - Direct Acquisition (Wave 1): WhatsApp (5,000 @ 5% = 250) + Clubs (3,500 @ 4% = 140) = 390 registrations.
   - Referral Amplification (Wave 2): 390 × 25% share rate × 1.2 sign-ups = 117 registrations.
   - Total Organic Base Case = 507 registrations (+7 buffer over 500 goal).
2. Pre-Committed Decision Triggers:
   - Day 4 Behind (<270): Deploy ₹800 Meta/Instagram paid ad targeting final-year engineers in tier-2/3 colleges (est. 40-50 regs @ ₹16-20 CPL), onboard 5 replacement connectors in zero-registration colleges.
   - Day 6 Behind: Switch all messaging to 24-hour urgency ("Workshop closes tonight"), direct connector outreach, and spot cash prizes.
   - Ahead of Pace: Protect the ₹800 cash reserve and reinvest effort into live workshop attendance reminders.""",

    # Slide 5 Notes
    """SLIDE 5: WORKSHOP LAUNCHPAD PRODUCT DEMONSTRATION
Key Talking Points:
1. Why Build a Tool? Growth engineers build systems, not just spreadsheets. Workshop Launchpad gives the campaign owner live visibility into pacing and attribution.
2. Link Studio: Generates unique collision-resistant codes, UTM campaign URLs, and multi-angle WhatsApp messaging templates.
3. Registration Demo Gateway: A real student registration experience that captures attribution, logs the lead, and provides the student their own viral invite link.
4. Live Dashboard: Real-time actual vs. target SVG pacing curve, daily ingestion breakdown, referral leaderboard, and automatic contingency trigger warnings.
5. Funnel Model: Live recalculation with K-factor viral multiplier and paid reserve toggle.
6. Architecture: Built with a lightweight zero-dependency Node.js REST API with file persistence, fully runnable standalone or client-server."""
]

for idx, slide in enumerate(prs.slides):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = speaker_notes[idx]

output_path = 'Growth Plan 500 Registrations in 7 Days (Updated).pptx'
prs.save(output_path)
print(f'Successfully updated presentation: {output_path}')
