import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation('original_backup/Growth Plan 500 Registrations in 7 Days.pptx')

for idx, slide in enumerate(prs.slides, 1):
    print(f'=== SLIDE {idx} ===')
    for sidx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            text = ' | '.join([p.text for p in shape.text_frame.paragraphs if p.text.strip()])
            if text:
                print(f'  Shape {sidx}: {text}')
        elif shape.has_table:
            print(f'  Table {sidx}: {len(shape.table.rows)} rows x {len(shape.table.columns)} cols')
            for row in shape.table.rows:
                cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
                print(f'    Row: {" | ".join(cells)}')
